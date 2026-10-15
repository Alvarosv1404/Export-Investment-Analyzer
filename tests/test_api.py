"""Tests de la API.

Se centran en el contrato, no en el formato: que los codigos HTTP sean los
correctos y que la respuesta sea JSON valido. Este ultimo punto importa mas de
lo que parece: un NaN en cualquier parte del payload hace que `JSONResponse`
lanze excepcion y el endpoint devuelva 500 sin explicar por que.
"""

from __future__ import annotations

import json
import math
from io import BytesIO
from urllib.parse import urlencode

import pytest
from openpyxl import load_workbook

from exportanalysis.api.dossier import generate_excel_dossier
from exportanalysis.api.main import DIST_DIR, _json_safe, app
from exportanalysis.config import load_catalog
from exportanalysis.pipeline.analyze import analyze_product

fastapi_testclient = pytest.importorskip("fastapi.testclient")
TestClient = fastapi_testclient.TestClient
FINANCIAL_QUERY = urlencode(
    {
        "capacity_kg_year": 100000,
        "year1_volume_kg": 50000,
        "annual_volume_growth_pct": 5,
        "purchase_cost_usd_per_kg": 2,
        "variable_cost_usd_per_kg": 0.5,
        "capex_usd": 250000,
        "opex_fixed_usd_year": 30000,
        "working_capital_usd": 12000,
        "discount_rate": 0.12,
    }
)


@pytest.fixture(scope="module")
def client():
    return TestClient(app)


def _assert_json_valid(payload) -> None:
    """Recorre el payload buscando valores que JSON no admite."""
    if isinstance(payload, float):
        assert math.isfinite(payload), f"NaN o Infinity en el payload: {payload}"
    elif isinstance(payload, dict):
        for value in payload.values():
            _assert_json_valid(value)
    elif isinstance(payload, list):
        for item in payload:
            _assert_json_valid(item)


class TestJsonSafe:
    def test_nan_a_null(self):
        assert _json_safe(float("nan")) is None
        assert _json_safe(float("inf")) is None
        assert _json_safe(float("-inf")) is None

    def test_finitos_intactos(self):
        assert _json_safe(1.5) == 1.5
        assert _json_safe(0.0) == 0.0
        assert _json_safe(None) is None
        assert _json_safe("texto") == "texto"
        assert _json_safe(True) is True

    def test_recorre_estructuras_anidadas(self):
        out = _json_safe(
            {"a": [float("nan"), {"b": float("inf")}], "c": {"d": 1.0}}
        )
        assert out == {"a": [None, {"b": None}], "c": {"d": 1.0}}

    def test_el_resultado_es_serializable(self):
        json.dumps(_json_safe({"x": float("nan")}))


class TestHealth:
    def test_health(self, client):
        response = client.get("/health")
        assert response.status_code == 200
        assert response.json() == {"status": "ok"}


class TestProducts:
    def test_catalogo(self, client):
        response = client.get("/api/products")
        assert response.status_code == 200
        data = response.json()
        assert data["reporter"] == 604  # Peru en M49
        assert len(data["products"]) >= 1
        slugs = {p["slug"] for p in data["products"]}
        assert "cafe_verde" in slugs
        for product in data["products"]:
            assert len(product["hs6"]) == 6
        assert any(p.get("industry") for p in data["products"]) is True

    def test_sin_parametros_obligatorios(self, client):
        assert client.get("/api/products").status_code == 200

    def test_productos_sin_supuestos_se_analizan_parcialmente(self, client):
        """Un producto descubierto en data/ puede no tener supuestos.

        El catalogo ya no exige supuestos: si faltan, el analisis entrega el
        mercado completo y marca lo que falta, en vez de fallar con 500.
        """
        for product in load_catalog().products:
            response = client.get(f"/api/analysis/{product.slug}")
            assert response.status_code == 200, product.slug
            data = response.json()
            assert data["investment"] is None
            assert data["assumptions_missing"]


class TestAnalysis:
    def test_slug_inexistente_da_404(self, client):
        response = client.get("/api/analysis/no_existe")
        assert response.status_code == 404
        assert "detail" in response.json()

    def test_analisis_es_json_valido(self, client):
        """Regresion: un NaN en la serie-rompe la serializacion y da 500."""
        response = client.get("/api/analysis/cafe_verde")
        assert response.status_code == 200
        data = response.json()  # si el body no fuera JSON valido, falla aqui
        _assert_json_valid(data)

    def test_estructura_del_analisis(self, client):
        data = client.get("/api/analysis/cafe_verde").json()
        for key in (
            "slug",
            "product",
            "market",
            "competitors",
            "headroom",
            "unit_economics",
            "landed_cost",
            "investment",
            "data_quality",
        ):
            assert key in data, f"falta la clave {key}"

    def test_el_mercado_trae_anios_con_data(self, client):
        data = client.get("/api/analysis/cafe_verde").json()
        assert data["market"]["available"] is True
        assert len(data["market"]["series"]) >= 1
        assert data["market"]["latest_fob_usd"] > 0

    def test_primer_anio_sin_yoy_uses_null(self, client):
        """El primer ano no tiene ano previo: debe ser null, no NaN ni 0."""
        data = client.get("/api/analysis/cafe_verde").json()
        assert data["market"]["series"][0]["value_yoy"] is None

    def test_analisis_usa_demanda_importadora_y_proyeccion_de_dos_anios(self, client):
        data = client.get("/api/analysis/cafe_verde").json()
        assert data["import_demand"]["available"] is True
        assert data["projection"]["method"] == "moving_average"
        assert len(data["projection"]["forecast"]) == 2

    def test_falta_volumen_o_crecimiento_no_rebota(self, client):
        query = (
            "capacity_kg_year=100000"
            "&purchase_cost_usd_per_kg=2"
            "&variable_cost_usd_per_kg=0.5"
            "&capex_usd=250000"
            "&opex_fixed_usd_year=30000"
            "&working_capital_usd=12000"
            "&discount_rate=0.12"
        )
        response = client.get(f"/api/analysis/cafe_verde?{query}")
        assert response.status_code == 200
        data = response.json()
        assert data["investment"] is None
        assert "year1_volume_kg" in data["assumptions_missing"]
        assert "annual_volume_growth_pct" in data["assumptions_missing"]

    def test_modelo_financiero_solo_se_calcula_con_supuestos_completos(self, client):
        data = client.get(f"/api/analysis/cafe_verde?{FINANCIAL_QUERY}").json()
        assert data["investment"] is not None
        assert data["assumptions_missing"] == []
        # El API completa las volatilidades con defaults (15/15/20), asi que el
        # Monte Carlo se calcula aunque el frontend no las pida explicitamente.
        monte_carlo = data["investment"]["monte_carlo"]
        assert data["investment"]["monte_carlo_missing"] == []
        assert monte_carlo is not None
        assert monte_carlo["assumptions"]["price_volatility_pct"] == 15.0
        assert monte_carlo["assumptions"]["variable_cost_volatility_pct"] == 15.0
        assert monte_carlo["assumptions"]["volume_volatility_pct"] == 20.0
        assert data["investment"]["inputs"]["year1_volume_kg"] == 50000
        assert data["investment"]["inputs"]["annual_volume_growth_pct"] == 5
        ramp_source = data["investment"]["ramp_source"]
        assert ramp_source["source"] == "usuario"
        assert ramp_source["reason"]

    def test_monte_carlo_se_omite_sin_volatilidades(self):
        """Sin volatilidades el Monte Carlo se omite y se listan las faltantes.

        El API las completa con defaults (test de arriba), pero el pipeline debe
        seguir avisando cuando de verdad no hay con que simular.
        """
        data = analyze_product(
            "cafe_verde",
            capacity_kg_year=100000,
            year1_volume_kg=50000,
            annual_volume_growth_pct=5,
            purchase_cost_usd_per_kg=2,
            variable_cost_usd_per_kg=0.5,
            capex_usd=250000,
            opex_fixed_usd_year=30000,
            working_capital_usd=12000,
            discount_rate=0.12,
        )
        assert data["investment"] is not None
        assert data["investment"]["monte_carlo"] is None
        assert data["investment"]["monte_carlo_missing"] == [
            "price_volatility_pct",
            "variable_cost_volatility_pct",
            "volume_volatility_pct",
        ]

    def test_monte_carlo_usa_volatilidades_ingresadas(self, client):
        query = (
            f"{FINANCIAL_QUERY}"
            "&price_volatility_pct=12"
            "&purchase_cost_volatility_pct=10"
            "&volume_volatility_pct=15"
        )
        response = client.get(f"/api/analysis/cafe_verde?{query}")
        assert response.status_code == 200
        monte_carlo = response.json()["investment"]["monte_carlo"]
        assert monte_carlo["assumptions"]["price_volatility_pct"] == 12
        assert monte_carlo["assumptions"]["variable_cost_volatility_pct"] == 10
        assert monte_carlo["assumptions"]["volume_volatility_pct"] == 15
        assert sum(row["count"] for row in monte_carlo["histogram"]) == monte_carlo["n_simulations"]

    def test_costo_variable_unico_incluye_compra_y_proceso(self, client):
        query = (
            "capacity_kg_year=100000&year1_volume_kg=50000"
            "&annual_volume_growth_pct=5&variable_cost_usd_per_kg=2.5"
            "&capex_usd=250000&opex_fixed_usd_year=30000&discount_rate=0.12"
        )
        data = client.get(f"/api/analysis/cafe_verde?{query}").json()
        investment = data["investment"]
        assert investment["inputs"]["purchase_cost_usd_per_kg"] == 0
        assert investment["inputs"]["variable_cost_usd_per_kg"] == 2.5
        first_year = investment["pnl"][0]
        assert first_year["variable_costs_usd"] == pytest.approx(
            first_year["volume_kg"] * 2.5
        )

    def test_volatilidad_fuera_de_rango_da_422(self, client):
        assert client.get("/api/analysis/cafe_verde?price_volatility_pct=101").status_code == 422
        assert client.get("/api/analysis/cafe_verde?volume_volatility_pct=-0.1").status_code == 422

    def test_target_share_explicito(self, client):
        data = client.get(f"/api/analysis/cafe_verde?target_share=0.12&{FINANCIAL_QUERY}").json()
        assert data["investment"]["inputs"]["capture_share"] == 0.12

    def test_target_share_fuera_de_rango_da_422(self, client):
        assert client.get("/api/analysis/cafe_verde?target_share=1.5").status_code == 422
        assert client.get("/api/analysis/cafe_verde?target_share=-0.1").status_code == 422

    def test_precio_explicito_se_usa_y_se_marca(self, client):
        data = client.get("/api/analysis/cafe_verde?price_usd_per_kg=6.5").json()
        assert data["pricing"]["price_usd_per_kg"] == 6.5
        assert data["data_quality"]["price_source"] == "ajustado por el usuario"

    def test_precio_por_defecto_es_calculado(self, client):
        data = client.get("/api/analysis/cafe_verde").json()
        # Con el precio SUNAT disponible, el origen debe ser el de SUNAT
        assert data["data_quality"]["price_source"] == "precio SUNAT (FOB / peso neto)"
        assert data["pricing"]["price_usd_per_kg"] > 0

    def test_precio_invalido_da_422(self, client):
        assert client.get("/api/analysis/cafe_verde?price_usd_per_kg=0").status_code == 422
        assert client.get("/api/analysis/cafe_verde?price_usd_per_kg=-1").status_code == 422

    def test_la_fuente_declarada_es_offline(self, client):
        data = client.get("/api/analysis/cafe_verde").json()
        assert data["data_quality"]["market_data_source"] == "Trade Map (Excel local)"


class TestDossier:
    def test_dossier_pendiente_no_inventa_resultados_financieros(self):
        stream = generate_excel_dossier(
            {
                "product": {"name": "Producto de prueba", "hs6": "123456"},
                "assumptions_missing": ["capex_usd"],
            }
        )
        workbook = load_workbook(BytesIO(stream.getvalue()), read_only=True)
        values = [
            cell.value
            for row in workbook["Resumen Ejecutivo"].iter_rows()
            for cell in row
        ]
        assert "Pendiente" in values
        assert any("capex_usd" in str(value) for value in values)
        assert "12.0%" not in values
        assert not any("Monte Carlo" in str(value) for value in values)

    def test_dossier_incluye_supuestos_y_resultados_monte_carlo(self):
        stream = generate_excel_dossier(
            {
                "product": {"name": "Producto de prueba", "hs6": "123456"},
                "investment": {
                    "summary": {"npv_usd": 1000},
                    "inputs": {"discount_rate": 0.12},
                    "monte_carlo": {
                        "assumptions": {
                            "price_volatility_pct": 12,
                            "variable_cost_volatility_pct": 10,
                            "purchase_cost_volatility_pct": 10,
                            "volume_volatility_pct": 15,
                        },
                        "n_simulations": 2000,
                        "prob_loss_pct": 22.5,
                        "npv_p10_usd": -500,
                        "npv_p90_usd": 2500,
                    },
                },
            }
        )
        workbook = load_workbook(BytesIO(stream.getvalue()), read_only=True)
        values = [
            cell.value
            for row in workbook["Resumen Ejecutivo"].iter_rows()
            for cell in row
        ]
        assert "Volatilidad del precio FOB" in values
        assert 12 in values
        assert "Probabilidad simulada de VAN negativo" in values

    def test_dossier_muestra_la_tasa_ingresada(self):
        stream = generate_excel_dossier(
            {
                "product": {"name": "Producto de prueba", "hs6": "123456"},
                "investment": {
                    "inputs": {"discount_rate": 0.14},
                    "summary": {
                        "npv_usd": 100,
                        "irr_pct": 18,
                        "payback_years": 3,
                        "verdict": {"label": "Favorable", "detail": "Caso de prueba"},
                    },
                },
            }
        )
        workbook = load_workbook(BytesIO(stream.getvalue()), read_only=True)
        values = [
            cell.value
            for row in workbook["Resumen Ejecutivo"].iter_rows()
            for cell in row
        ]
        assert "14.0%" in values

    def test_dossier_omite_metricas_no_respaldadas_y_aranceles(self):
        stream = generate_excel_dossier(
            {
                "product": {"name": "Café verde", "hs6": "090111"},
                "unit_economics": {
                    "gross_margin_usd_per_kg": 1.5,
                    "gross_margin_pct": 0.2,
                    "max_purchase_price_usd_per_kg": 2,
                    "breakeven_fob_price_usd_per_kg": 3,
                },
                "tariffs": [{"country": "Destino", "rate": 10}],
                "import_demand": {
                    "destinations": [
                        {
                            "country": "Destino",
                            "year": 2025,
                            "import_value_usd": 1000,
                            "import_quantity_kg": 500,
                            "unit_value_usd_per_kg": 2,
                        }
                    ]
                },
            }
        )
        workbook = load_workbook(BytesIO(stream.getvalue()), read_only=True)
        summary_values = [
            cell.value
            for row in workbook["Resumen Ejecutivo"].iter_rows()
            for cell in row
        ]
        assert "Margen Bruto Unitario" not in summary_values
        assert "Precio Máx. Tolerable en Chacra" not in summary_values
        assert "Precio FOB de Equilibrio" not in summary_values
        assert "Aranceles en Destino" not in workbook.sheetnames
        assert "Demanda Importadora" in workbook.sheetnames


class TestComparison:
    def test_comparacion_es_json_valido(self, client):
        response = client.get("/api/comparison?slugs=cafe_verde,uva_fresca")
        assert response.status_code == 200
        data = response.json()
        _assert_json_valid(data)
        assert data["available_count"] >= 2

    def test_comparacion_sin_slugs_usa_todos(self, client):
        data = client.get("/api/comparison").json()
        assert data["available_count"] >= 1


class TestHtml:
    def test_raiz_renderiza(self, client):
        """Sirve el build de Vite si existe; si no, el HTML de Jinja."""
        response = client.get("/")
        assert response.status_code == 200
        assert "text/html" in response.headers["content-type"]

    def test_build_de_vite_devuelve_el_shell(self, client):
        """Con dist/ presente, / devuelve el shell y los datos los pide el JS.

        El shell no lleva el nombre del producto: el estado vive en la URL y
        `frontend/js/main.js` lo lee para llamar a /api/analysis/...
        """
        if not (DIST_DIR / "index.html").exists():
            pytest.skip("sin build de Vite; corre `npm run build`")
        response = client.get("/?slug=cafe_verde")
        assert response.status_code == 200
        assert 'id="app"' in response.text
        assert "assets/index" in response.text

    def test_ssr_fallback_renderiza_el_producto(self, client, monkeypatch):
        """Con FORCE_SSR=1 el HTML llega completo desde el servidor.

        El render del lado del servidor es un reporte imprimible de solo lectura:
        no lleva formularios de supuestos (eso lo arma el JS del navegador), pero
        si todas las secciones —incluida la parte financiera— ya resueltas.
        """
        monkeypatch.setenv("FORCE_SSR", "1")
        response = client.get(f"/?slug=cafe_verde&{FINANCIAL_QUERY}")
        assert response.status_code == 200
        assert "Café verde" in response.text
        for heading in (
            "1. Economía unitaria",
            "2. Mercado",
            "3. Principales destinos",
            "4. Competencia regional",
            "5. Espacio de mercado",
            "6. Aranceles y costo puesto en destino",
            "7. Decisión de inversión",
            "8. Sensibilidad",
            "Procedencia de los datos",
        ):
            assert heading in response.text
        # Las cifras financieras se resuelven en el servidor, no con JS.
        assert "Margen bruto" in response.text
        assert "Precio de equilibrio" in response.text
        assert "TIR" in response.text
        # Los graficos se dibujan sobre <canvas> con la data embebida para Chart.js.
        assert "<canvas" in response.text
        assert "const R = " in response.text
        # Sin JS no hay formularios: los supuestos se editan en el navegador.
        assert "<input" not in response.text
        assert "<select" not in response.text

    def test_ssr_fallback_toma_el_primer_producto(self, client, monkeypatch):
        monkeypatch.setenv("FORCE_SSR", "1")
        response = client.get("/")
        assert response.status_code == 200
        assert "Café verde" in response.text

    def test_ssr_fallback_avisa_si_el_producto_no_existe(self, client, monkeypatch):
        """Un slug malo no debe ser un 500: es un mensaje en la pagina."""
        monkeypatch.setenv("FORCE_SSR", "1")
        response = client.get("/?slug=no_existe")
        assert response.status_code == 200

    def test_sirve_alguna_hoja_de_estilos(self, client):
        """Con build hay /assets/index.css; sin build, /static/app.css."""
        built = client.get("/assets/index.css")
        legacy = client.get("/static/app.css")
        assert built.status_code == 200 or legacy.status_code == 200
